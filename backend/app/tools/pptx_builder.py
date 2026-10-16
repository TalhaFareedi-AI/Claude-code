"""PPT builder (section 6, page 25-26): the AI writes slide data in JSON, our code builds the file from templates.
At most 5 bullets per slide, at most 12 words each. Charts become real, editable PPT charts.
Themes: Academic, Business, Minimal (plus clean and night)."""
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches, Pt

from .. import charts

THEMES = {
    "clean": {"bg": "FFFFFF", "title": "1F3A5F", "text": "222222", "accent": "2E7D6B", "bar": "E6F4F1"},
    "academic": {"bg": "FCFBF7", "title": "1B365D", "text": "1E293B", "accent": "8B2635", "bar": "EFECE6"},
    "business": {"bg": "FFFFFF", "title": "0F172A", "text": "334155", "accent": "2563EB", "bar": "EFF6FF"},
    "minimal": {"bg": "FAFAFA", "title": "18181B", "text": "3F3F46", "accent": "2A7F6F", "bar": "F4F4F5"},
    "night": {"bg": "14213D", "title": "FCA311", "text": "F1F1F1", "accent": "6FCF97", "bar": "1E2E52"},
}
MAX_BULLETS, MAX_WORDS = 5, 12


def clamp_slide(s: dict) -> dict:
    """Code enforces the limits so the checker does not need an AI call."""
    bullets = []
    for b in (s.get("bullets") or [])[:MAX_BULLETS]:
        words = str(b).split()
        bullets.append(" ".join(words[:MAX_WORDS]))
    return {**s, "title": str(s.get("title", ""))[:80], "bullets": bullets}


def check_slides(slides: list[dict]) -> list[str]:
    problems = []
    for i, s in enumerate(slides, 1):
        if not s.get("title"):
            problems.append(f"slide {i} has no title")
        if not s.get("bullets") and not s.get("chart"):
            problems.append(f"slide {i} is empty")
    return problems


def build(deck: dict, out: Path, theme: str | None = None) -> Path:
    theme_key = (theme or deck.get("theme") or "clean").lower().strip()
    th = THEMES.get(theme_key, THEMES["clean"])
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]
    rgb = lambda h: RGBColor.from_string(h)
    slides = [clamp_slide(s) for s in deck.get("slides", [])]
    total = len(slides)
    deck_title = str(deck.get("title") or "Presentation")[:70]

    for idx, s in enumerate(slides):
        slide = prs.slides.add_slide(blank)
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = rgb(th["bg"])

        # Top title textbox
        tb = slide.shapes.add_textbox(Inches(0.7), Inches(0.4), Inches(11.9), Inches(1.1))
        p = tb.text_frame.paragraphs[0]
        p.text = s["title"]
        p.font.size, p.font.bold, p.font.color.rgb = Pt(38 if idx else 44), True, rgb(th["title"])

        has_chart = bool(s.get("chart"))
        body_w = Inches(5.6) if has_chart else Inches(11.9)
        if s["bullets"]:
            body = slide.shapes.add_textbox(Inches(0.9), Inches(1.8), body_w, Inches(4.8))
            tf = body.text_frame
            tf.word_wrap = True
            for n, b in enumerate(s["bullets"]):
                para = tf.paragraphs[0] if n == 0 else tf.add_paragraph()
                para.text = "• " + b
                para.font.size, para.font.color.rgb = Pt(24), rgb(th["text"])
                para.space_after = Pt(12)
        if has_chart:
            try:
                _add_chart(slide, s["chart"], Inches(6.7), Inches(1.8), Inches(6.0), Inches(4.8))
            except Exception:
                pass
        elif s.get("visual", {}).get("type") in ("diagram", "image") and s["visual"].get("prompt"):
            box = slide.shapes.add_textbox(Inches(0.9), Inches(6.2), Inches(11.5), Inches(0.6))
            box.text_frame.text = "Visual: " + str(s["visual"]["prompt"])[:120]
            box.text_frame.paragraphs[0].font.size = Pt(14)
            box.text_frame.paragraphs[0].font.italic = True
            box.text_frame.paragraphs[0].font.color.rgb = rgb(th["accent"])

        # Footer with deck title and slide number
        ftr = slide.shapes.add_textbox(Inches(0.7), Inches(6.9), Inches(11.9), Inches(0.4))
        fp = ftr.text_frame.paragraphs[0]
        fp.text = f"{deck_title}  ·  Slide {idx + 1} of {total}"
        fp.font.size = Pt(11)
        fp.font.color.rgb = rgb(th["accent"])

        if s.get("notes"):
            slide.notes_slide.notes_text_frame.text = str(s["notes"])[:600]
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out)
    return out


def _add_chart(slide, recipe: dict, x, y, w, h) -> None:
    rec = charts.validate(recipe)
    data = charts.compute(rec)
    cd = CategoryChartData()
    if "xs" in data:
        step = max(1, len(data["xs"]) // 20)
        idxs = list(range(0, len(data["xs"]), step))
        cd.categories = [f"{data['xs'][i]:.3g}" for i in idxs]
        for s in data["series"]:
            cd.add_series(s["name"], [s["values"][i] for i in idxs])
        kind = XL_CHART_TYPE.LINE if rec["chart"] in ("line", "scatter") else XL_CHART_TYPE.AREA
    else:
        cd.categories = data["labels"]
        for s in data["series"]:
            cd.add_series(s["name"], s["values"])
        kind = XL_CHART_TYPE.PIE if rec["chart"] == "pie" else XL_CHART_TYPE.COLUMN_CLUSTERED
    gf = slide.shapes.add_chart(kind, x, y, w, h, cd)
    gf.chart.has_title = True
    gf.chart.chart_title.text_frame.text = rec["title"]
